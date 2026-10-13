"""Document ingestion, lightweight RAG retrieval, and student progress storage."""

from __future__ import annotations

import hashlib

import io

import ipaddress

import math

import re

import socket

import sqlite3

import tempfile

from collections import Counter

from contextlib import contextmanager

from datetime import datetime, timezone

from difflib import SequenceMatcher, get_close_matches

from html.parser import HTMLParser

from pathlib import Path

from urllib.parse import urljoin, urlsplit, urlunsplit

import requests

BASE_DIR = Path(__file__).resolve().parent

DB_PATH = BASE_DIR / "data" / "studymate.sqlite3"

_STOP_WORDS = {

    "a", "an", "and", "are", "as", "at", "be", "been", "being", "but", "by",

    "for", "from", "had", "has", "have", "he", "her", "hers", "him", "his",

    "how", "i", "if", "in", "into", "is", "it", "its", "me", "my", "of", "on",

    "or", "our", "ours", "she", "that", "the", "their", "them", "there",

    "these", "they", "this", "those", "to", "us", "was", "we", "were", "what",

    "when", "where", "which", "who", "why", "will", "with", "you", "your",

}

_TOPIC_STOP_WORDS = _STOP_WORDS | {

    "about", "also", "can", "could", "does", "each", "first", "get", "like",

    "make", "made", "many", "more", "most", "much", "new", "one", "other",

    "often", "some", "such", "than", "then", "use", "used", "using", "very",

    "way", "well", "work", "would",

}

@contextmanager

def db():

    DB_PATH.parent.mkdir(parents=True, exist_ok=True)

    con = sqlite3.connect(DB_PATH)

    con.row_factory = sqlite3.Row

    con.execute("PRAGMA foreign_keys=ON")

    try:

        yield con

        con.commit()

    finally:

        con.close()

def init_db():

    with db() as con:

        con.executescript("""

            CREATE TABLE IF NOT EXISTS documents (

                id INTEGER PRIMARY KEY,

                name TEXT NOT NULL,

                sha256 TEXT NOT NULL UNIQUE,

                kind TEXT NOT NULL,

                uploaded_at TEXT NOT NULL

            );

            CREATE TABLE IF NOT EXISTS learning_topics (

                id INTEGER PRIMARY KEY,

                document_id INTEGER NOT NULL

                    REFERENCES documents(id) ON DELETE CASCADE,

                topic_order INTEGER NOT NULL,

                topic_name TEXT NOT NULL,

                completed INTEGER NOT NULL DEFAULT 0,

                UNIQUE(document_id, topic_order)

            );

            CREATE TABLE IF NOT EXISTS chunks (

                id INTEGER PRIMARY KEY,

                document_id INTEGER NOT NULL REFERENCES documents(id) ON DELETE CASCADE,

                source TEXT NOT NULL,

                content TEXT NOT NULL

            );

            CREATE TABLE IF NOT EXISTS attempts (

                id INTEGER PRIMARY KEY,

                topic TEXT NOT NULL,

                question TEXT NOT NULL,

                chosen TEXT NOT NULL,

                correct INTEGER NOT NULL,

                difficulty TEXT NOT NULL,

                created_at TEXT NOT NULL

            );

        """)

def chunks_from_text(text: str, source: str, size=165, overlap=28):

    words = re.sub(r"\s+", " ", text).strip().split()

    if not words:

        return []

    step = max(1, size - overlap)

    return [(source, " ".join(words[i:i+size])) for i in range(0, len(words), step)]

class _ArticleTextParser(HTMLParser):

    def __init__(self):

        super().__init__(convert_charrefs=True)

        self.title_parts = []

        self.text_parts = []

        self._in_title = False

        self._skip_depth = 0

        self._skip_tags = {"script", "style", "noscript", "svg", "nav", "footer", "header"}

    def handle_starttag(self, tag, attrs):

        if tag == "title":

            self._in_title = True

        if tag in self._skip_tags:

            self._skip_depth += 1

        if tag in {"p", "br", "div", "li", "h1", "h2", "h3", "h4", "tr"}:

            self.text_parts.append("\n")

    def handle_endtag(self, tag):

        if tag == "title":

            self._in_title = False

        if tag in self._skip_tags and self._skip_depth:

            self._skip_depth -= 1

        if tag in {"p", "br", "div", "li", "h1", "h2", "h3", "h4", "tr"}:

            self.text_parts.append("\n")

    def handle_data(self, data):

        if self._in_title:

            self.title_parts.append(data)

        if not self._skip_depth:

            self.text_parts.append(data)

def _validated_public_url(url: str) -> str:

    if len(url) > 2048:

        raise ValueError("The URL is too long.")

    parsed = urlsplit(url.strip())

    if parsed.scheme.lower() not in {"http", "https"} or not parsed.hostname:

        raise ValueError("Enter a full public http:// or https:// URL.")

    if parsed.username or parsed.password:

        raise ValueError("URLs with embedded usernames or passwords are not supported.")

    if parsed.port not in (None, 80, 443):

        raise ValueError("Only standard HTTP and HTTPS ports are supported.")

    host = parsed.hostname.rstrip(".").encode("idna").decode("ascii")

    try:

        addresses = {result[4][0] for result in socket.getaddrinfo(host, parsed.port or (443 if parsed.scheme == "https" else 80))}

    except socket.gaierror as exc:

        raise ValueError("The URL host could not be resolved.") from exc

    if not addresses or any(not ipaddress.ip_address(address).is_global for address in addresses):

        raise ValueError("Only publicly reachable website URLs can be imported.")

    return urlunsplit((parsed.scheme.lower(), parsed.netloc, parsed.path or "/", parsed.query, ""))

def _fetch_web_page(url: str) -> tuple[str, bytes]:

    current_url = url

    for _ in range(6):

        current_url = _validated_public_url(current_url)

        try:

            response = requests.get(

                current_url,

                headers={"User-Agent": "StudyMateAI/1.0 (educational page importer)"},

                timeout=(5, 15),

                stream=True,

                allow_redirects=False,

            )

        except requests.RequestException as exc:

            raise RuntimeError(f"Could not fetch the web page: {exc}") from exc

        with response:

            if response.is_redirect:

                location = response.headers.get("Location")

                if not location:

                    raise RuntimeError("The website returned a redirect without a destination.")

                current_url = urljoin(current_url, location)

                continue

            if response.status_code != 200:

                raise RuntimeError(f"The website returned HTTP {response.status_code}.")

            content_type = response.headers.get("Content-Type", "").split(";", 1)[0].strip().lower()

            if content_type not in {"text/html", "application/xhtml+xml"}:

                raise ValueError("That URL does not return an HTML web page.")

            content = bytearray()

            for chunk in response.iter_content(chunk_size=65536):

                content.extend(chunk)

                if len(content) > 2_000_000:

                    raise ValueError("The web page is larger than the 2 MB import limit.")

            return current_url, bytes(content)

    raise RuntimeError("The website redirected too many times.")

def add_url_document(url: str):

    """Fetch and index readable text from a publicly accessible web page."""

    final_url, data = _fetch_web_page(url)

    parser = _ArticleTextParser()

    try:

        parser.feed(data.decode("utf-8", errors="replace"))

    except Exception as exc:

        raise ValueError("Could not read the page's HTML.") from exc

    title = re.sub(r"\s+", " ", " ".join(parser.title_parts)).strip()

    text = re.sub(r"\s+", " ", " ".join(parser.text_parts)).strip()

    if len(text) < 40:

        raise ValueError("The page did not contain enough readable text to make a study source.")

    host = urlsplit(final_url).hostname or "website"

    name = (title or f"Lesson from {host}")[:200]

    if title:

        text = f"{title}. {text}"

    digest = hashlib.sha256((final_url + "\n" + text).encode("utf-8")).hexdigest()

    with db() as con:

        prior = con.execute("SELECT id FROM documents WHERE sha256=?", (digest,)).fetchone()

    if prior:

        return {"added": False, "chunks": 0, "message": "This page is already in your library"}

    content = chunks_from_text(text, f"Web page · {final_url}")

    now = datetime.now(timezone.utc).isoformat()

    with db() as con:

        cur = con.execute(

            "INSERT INTO documents(name, sha256, kind, uploaded_at) VALUES (?, ?, ?, ?)",

            (name, digest, "url", now),

        )

        con.executemany(

            "INSERT INTO chunks(document_id, source, content) VALUES (?, ?, ?)",

            [(cur.lastrowid, label, passage) for label, passage in content],

        )

    return {"added": True, "chunks": len(content), "message": f"Indexed {len(content)} source passages"}

def extract_file(filename: str, data: bytes):

    """Return [(source locator, extracted text), ...]. Do not trust filename as a path."""

    suffix = Path(filename).suffix.lower()

    if suffix == ".pdf":

        import fitz

        with fitz.open(stream=data, filetype="pdf") as doc:

            return [chunk for page_num, page in enumerate(doc, 1)

                    for chunk in chunks_from_text(page.get_text(), f"{filename} · page {page_num}")]

    if suffix == ".pptx":

        from pptx import Presentation

        presentation = Presentation(io.BytesIO(data))

        result = []

        for n, slide in enumerate(presentation.slides, 1):

            pieces = []

            for shape in slide.shapes:

                if shape.has_text_frame:

                    pieces.append(shape.text)

                if shape.has_table:

                    for row in shape.table.rows:

                        pieces.append(" | ".join(cell.text for cell in row.cells))

            result += chunks_from_text("\n".join(pieces), f"{filename} · slide {n}")

        return result

    if suffix in (".txt", ".md"):

        return chunks_from_text(data.decode("utf-8", errors="replace"), f"{filename} · text")

    if suffix in (".mp4", ".mp3", ".m4a", ".wav"):

        return transcribe_media(filename, data)

    raise ValueError(f"Unsupported format: {suffix}")

def fmt_time(seconds):

    seconds = max(0, int(seconds))

    return f"{seconds // 3600:02d}:{(seconds // 60) % 60:02d}:{seconds % 60:02d}"

def transcribe_media(filename: str, data: bytes):

    try:

        from faster_whisper import WhisperModel

    except ImportError as exc:

        raise RuntimeError("Video/audio processing requires: pip install -r requirements-video.txt") from exc

    suffix = Path(filename).suffix.lower()

    with tempfile.TemporaryDirectory() as temp_dir:

        media_path = Path(temp_dir) / ("upload" + suffix)

        media_path.write_bytes(data)

        model = WhisperModel("base", device="cpu", compute_type="int8")

        segments, _ = model.transcribe(str(media_path), vad_filter=True)

        result = []

        buffer = []

        start = 0

        for seg in segments:

            if not buffer:

                start = seg.start

            buffer.append(seg.text.strip())

            if sum(len(item.split()) for item in buffer) >= 120:

                result += chunks_from_text(" ".join(buffer), f"{filename} · {fmt_time(start)}")

                buffer = []

        if buffer:

            result += chunks_from_text(" ".join(buffer), f"{filename} · {fmt_time(start)}")

        return result

def add_document(filename: str, data: bytes):

    digest = hashlib.sha256(data).hexdigest()

    with db() as con:

        prior = con.execute("SELECT id FROM documents WHERE sha256=?", (digest,)).fetchone()

    if prior:

        return {"added": False, "chunks": 0, "message": "Already in your library"}

    content = extract_file(filename, data)

    if not content:

        raise ValueError("No readable text was found. Scanned PDFs require OCR (not included).")

    kind = Path(filename).suffix.lower().lstrip(".")

    now = datetime.now(timezone.utc).isoformat()

    with db() as con:

        cur = con.execute("INSERT INTO documents(name, sha256, kind, uploaded_at) VALUES (?, ?, ?, ?)",

                          (Path(filename).name, digest, kind, now))

        con.executemany("INSERT INTO chunks(document_id, source, content) VALUES (?, ?, ?)",

                        [(cur.lastrowid, label, text) for label, text in content])

    return {"added": True, "chunks": len(content), "message": f"Indexed {len(content)} source passages"}

def add_diagram_document(filename: str, image_data: bytes, description: str):

    """Index an image's vision-generated description with an image citation."""

    if not image_data:

        raise ValueError("The uploaded image is empty.")

    if not description.strip():

        raise ValueError("The diagram description is empty.")

    digest = hashlib.sha256(image_data).hexdigest()

    with db() as con:

        prior = con.execute("SELECT id FROM documents WHERE sha256=?", (digest,)).fetchone()

    if prior:

        return {"added": False, "chunks": 0, "message": "This image is already in your library"}

    content = chunks_from_text(description.strip(), f"Diagram image · {Path(filename).name}")

    now = datetime.now(timezone.utc).isoformat()

    with db() as con:

        cur = con.execute(

            "INSERT INTO documents(name, sha256, kind, uploaded_at) VALUES (?, ?, ?, ?)",

            (Path(filename).name, digest, "image", now),

        )

        con.executemany(

            "INSERT INTO chunks(document_id, source, content) VALUES (?, ?, ?)",

            [(cur.lastrowid, label, passage) for label, passage in content],

        )

    return {"added": True, "chunks": len(content), "message": f"Indexed {len(content)} diagram passages"}

def get_documents():

    with db() as con:

        return [dict(row) for row in con.execute("""

            SELECT d.id, d.name, d.kind, d.uploaded_at, COUNT(c.id) AS passages

            FROM documents d LEFT JOIN chunks c ON c.document_id=d.id

            GROUP BY d.id ORDER BY d.id DESC""").fetchall()]

def get_chunks(document_id: int | None = None):

    with db() as con:

        query = """

            SELECT c.id, c.document_id, c.source, c.content, d.name FROM chunks c

            JOIN documents d ON c.document_id=d.id

        """

        if document_id is None:

            rows = con.execute(query + " ORDER BY c.id").fetchall()

        else:

            rows = con.execute(query + " WHERE c.document_id=? ORDER BY c.id", (document_id,)).fetchall()

        return [dict(row) for row in rows]

def web_source_url(source: str) -> str | None:

    """Return the URL from a web citation, if the source label contains one."""

    prefix = "Web page · "

    if not source.startswith(prefix):

        return None

    url = source[len(prefix):].strip()

    parsed = urlsplit(url)

    if parsed.scheme not in {"http", "https"} or not parsed.hostname or parsed.username or parsed.password:

        return None

    return url

def get_document_topics(document_id: int, limit: int = 8):
    """Offer distinct, source-derived topics when no AI learning path exists.

    This is an offline fallback: the AI-generated learning path is more accurate
    for long or complex material.
    """
    chunks = get_chunks(document_id)
    if not chunks:
        return []
    with db() as con:
        row = con.execute("SELECT name FROM documents WHERE id=?", (document_id,)).fetchone()
    title = row["name"] if row else ""
    title_words = set(re.findall(r"[a-z]{3,}", title.lower()))
    generic = _TOPIC_STOP_WORDS | {
        "tutorial", "course", "chapter", "lesson", "material", "materials",
        "website", "index", "examples", "example", "section", "click",
        "home", "page", "learn", "learning", "introduction", "contents",
    }
    # Most learning sites repeat their course title on every page. Do not offer
    # that as the only topic.
    dominant = title_words & {"python", "tutorial", "guide", "course", "learning", "introduction"}
    ignored = generic | dominant
    corpus = " ".join(chunk["content"] for chunk in chunks)
    tokens = re.findall(r"(?u)\b[a-z][a-z0-9_'-]{2,}\b", corpus.lower())
    if not tokens:
        return []

    # For teaching materials that explicitly enumerate concepts in a known
    # syllabus, retain the order in which the actual text mentions them.
    known = [
        ("Variables", r"\bvariables?\b"),
        ("Data Types", r"\bdata types?\b"),
        ("Strings", r"\bstrings?\b"),
        ("Lists", r"\blists?\b"),
        ("Tuples", r"\btuples?\b"),
        ("Sets", r"\bsets?\b"),
        ("Dictionaries", r"\bdictionar(?:y|ies)\b"),
        ("Conditions", r"\b(?:if statements?|conditions?)\b"),
        ("Loops", r"\bloops?\b"),
        ("Functions", r"\bfunctions?\b"),
        ("Classes and Objects", r"\bclasses?\b.{0,40}\bobjects?\b"),
        ("File Handling", r"\bfile handling\b"),
        ("Exceptions", r"\bexceptions?\b"),
        ("Modules", r"\bmodules?\b"),
    ] if "python" in title.lower() else []
    named = []
    for label, pattern in known:
        match = re.search(pattern, corpus, re.I)
        if match:
            named.append((match.start(), label))
    named.sort()
    selections = [name for _, name in named]
    if len(selections) >= 3:
        return selections[:limit]

    # For arbitrary lessons, prefer sentence subjects / heading-like labels,
    # not accidental word pairs formed across sentence boundaries.
    verbs = {
        "is", "are", "was", "were", "uses", "use", "helps", "help",
        "allows", "allow", "contains", "contain", "provides", "provide",
        "creates", "create", "produces", "produce", "releases", "release",
        "stores", "store", "defines", "define", "captures", "capture",
        "needs", "need", "explains", "explain", "involves", "involve",
        "can", "may", "includes", "include", "supports", "support",
        "refers", "means", "becomes", "begins", "works", "consists",
        "happens", "requires", "form", "forms", "occurs", "occurring",
    }
    sentences = re.split(r"(?<=[.!?])\s+|[\n\r]+", corpus)
    for sentence in sentences:
        # Keep candidate labels grounded in the start of an actual sentence.
        sentence = re.sub(r"^[^A-Za-z]+", "", sentence).strip()
        words = re.findall(r"[A-Za-z][A-Za-z0-9_'-]*", sentence)
        if not words or len(words) < 3:
            continue
        subject = []
        for word in words[:6]:
            w = word.lower()
            if w in verbs:
                break
            if w in _STOP_WORDS and not subject:
                continue
            if w in ignored and not subject:
                continue
            subject.append(word)
            if len(subject) == 3:
                break
        if not subject:
            continue
        # Example "Cellular respiration uses glucose" -> "Cellular Respiration"
        label = " ".join(subject).title()
        if (len(label) < 5 or label.lower() in {"plants", "people", "therefore", "however",
                                                 "different", "important", "example"}):
            continue
        if any(label.casefold() in name.casefold() or name.casefold() in label.casefold()
               for name in selections):
            continue
        selections.append(label)
        if len(selections) >= limit:
            return selections[:limit]

    # Backstop for material without natural sentence breaks.
    candidates = Counter()
    for chunk in chunks:
        for sentence in re.split(r"(?<=[.!?])\s+|[\n\r]+", chunk["content"]):
            words = re.findall(r"(?u)\b[a-z][a-z0-9_'-]{2,}\b", sentence.lower())
            for n in (2, 3):
                for i in range(len(words) - n + 1):
                    seq = words[i:i+n]
                    if any(w in ignored or w in verbs for w in seq):
                        continue
                    candidates[" ".join(seq)] += 1
    for term, _ in candidates.most_common(limit * 10):
        if len(selections) >= limit:
            break
        label = term.title()
        if any(term in item.lower() or item.lower() in term for item in selections):
            continue
        selections.append(label)
    if not selections:
        counts = Counter(w for w in tokens if w not in ignored and w not in verbs)
        selections = [word.title() for word, _ in counts.most_common(limit)]
    return selections[:limit]


def ensure_learning_path(document_id: int):
    """Create an offline source-derived learning path if one isn't saved yet."""
    saved = get_learning_path(document_id)
    if saved:
        return saved
    topics = get_document_topics(document_id)
    if topics:
        save_learning_path(document_id, topics)
    return get_learning_path(document_id)

def retrieve(query: str, limit=4, *, document_id: int | None = None):

    rows = get_chunks(document_id)

    if not rows or not query.strip():

        return []

    def terms(text, vocabulary=None):

        words = [word for word in re.findall(r"(?u)\b\w\w+\b", text.lower())

                 if word not in _STOP_WORDS]

        if vocabulary:

            corrected = []

            for word in words:

                if word in vocabulary or len(word) < 5:

                    corrected.append(word)

                    continue

                candidates = get_close_matches(word, vocabulary, n=2, cutoff=0.82)

                if not candidates:

                    corrected.append(word)

                    continue

                best_score = SequenceMatcher(None, word, candidates[0]).ratio()

                second_score = SequenceMatcher(None, word, candidates[1]).ratio() if len(candidates) > 1 else 0

                corrected.append(candidates[0] if best_score - second_score >= 0.05 else word)

            words = corrected

        return Counter(words + [f"{left} {right}" for left, right in zip(words, words[1:])])

    document_terms = [terms(row["content"]) for row in rows]

    vocabulary = {term for frequencies in document_terms for term in frequencies if " " not in term}

    query_terms = terms(query, vocabulary)

    if not query_terms:

        return []

    document_frequency = Counter(

        term for frequencies in document_terms for term in frequencies

    )

    document_count = len(rows)

    def vector(frequencies):

        weighted = {

            term: (1 + math.log(count)) *

            (math.log((1 + document_count) / (1 + document_frequency[term])) + 1)

            for term, count in frequencies.items()

        }

        norm = math.sqrt(sum(weight * weight for weight in weighted.values()))

        return {term: weight / norm for term, weight in weighted.items()} if norm else {}

    query_vector = vector(query_terms)

    ranked = []

    for index, frequencies in enumerate(document_terms):

        document_vector = vector(frequencies)

        score = sum(weight * document_vector.get(term, 0) for term, weight in query_vector.items())

        if score > 0.005:

            ranked.append((score, index))

    ranked.sort(reverse=True)

    return [

        {**rows[index], "score": score, "marker": f"S{rank + 1}"}

        for rank, (score, index) in enumerate(ranked[:limit])

    ]

def save_attempt(topic: str, question: str, chosen: str, correct: bool, difficulty: str):

    with db() as con:

        con.execute("INSERT INTO attempts(topic, question, chosen, correct, difficulty, created_at) VALUES (?, ?, ?, ?, ?, ?)",

                    (topic.strip().lower(), question, chosen, int(correct), difficulty,

                     datetime.now(timezone.utc).isoformat()))

def mastery(topic: str):

    with db() as con:

        rows = con.execute("SELECT correct, difficulty FROM attempts WHERE topic=? ORDER BY id DESC LIMIT 12",

                           (topic.strip().lower(),)).fetchall()

    if not rows:

        return 45  # baseline for a new topic, not a measured mastery score

    weights = {"easy": 0.8, "medium": 1.0, "hard": 1.2}

    weighted_correct = sum(row["correct"] * weights.get(row["difficulty"], 1) for row in rows)

    weighted_total = sum(weights.get(row["difficulty"], 1) for row in rows)

    return round(100 * weighted_correct / weighted_total) if weighted_total else 45

def difficulty_for(topic: str):

    score = mastery(topic)

    return "easy" if score < 40 else "medium" if score < 75 else "hard"

def topic_stats():

    with db() as con:

        topics = [row[0] for row in con.execute("SELECT DISTINCT topic FROM attempts ORDER BY topic")]

        counts = {row[0]: row[1] for row in con.execute("SELECT topic, COUNT(*) FROM attempts GROUP BY topic")}

    return [{"topic": topic, "mastery": mastery(topic), "attempts": counts[topic]} for topic in topics]

def activity():

    with db() as con:

        return [dict(row) for row in con.execute("SELECT created_at, topic, correct FROM attempts ORDER BY id DESC LIMIT 50")]

# ---------- SMART LEARNING PATH ----------

def save_learning_path(document_id: int, topics: list[str], *, replace: bool = False):
    """Store source topics, optionally replacing topics while preserving matches."""
    cleaned = []
    seen = set()
    for topic in topics:
        name = " ".join(str(topic).split())[:100]
        if name and name.casefold() not in seen:
            cleaned.append(name)
            seen.add(name.casefold())
    if not cleaned:
        raise ValueError("No learning topics to save.")
    with db() as con:
        rows = con.execute(
            "SELECT topic_name, completed FROM learning_topics WHERE document_id=?",
            (document_id,)
        ).fetchall()
        if rows and not replace:
            return False
        completed = {r["topic_name"].casefold() for r in rows if r["completed"]}
        if rows:
            con.execute("DELETE FROM learning_topics WHERE document_id=?", (document_id,))
        con.executemany("""
            INSERT INTO learning_topics(document_id, topic_order, topic_name, completed)
            VALUES (?, ?, ?, ?)
        """, [(document_id, idx, name, int(name.casefold() in completed))
               for idx, name in enumerate(cleaned, 1)])
    return True

def get_learning_path(document_id: int):

    """Retrieve topics in their learning order."""

    with db() as con:

        rows = con.execute(

            """

            SELECT id, topic_order, topic_name, completed

            FROM learning_topics

            WHERE document_id=?

            ORDER BY topic_order

            """,

            (document_id,)

        ).fetchall()

    return [dict(row) for row in rows]

def complete_learning_topic(document_id: int, topic_name: str) -> bool:
    """Persist completion after a passing quiz, including a missed-question retry."""
    with db() as con:
        cursor = con.execute("""
            UPDATE learning_topics SET completed=1
            WHERE document_id=? AND lower(topic_name)=lower(?)
        """, (document_id, topic_name.strip()))
        return cursor.rowcount > 0
